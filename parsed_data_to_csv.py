import json
import codecs
from typing import Tuple, List

from docx import Document
from docx.table import Table

import tables as tabs

RESULTS = json.load(
    codecs.open("results.json", "r+", encoding="utf-8")
)  # open the file
CODES = json.load(open("codes.json", "r+", encoding="utf-8"))  # open the file
DOC = Document()

tab1 = tabs.Table1()
tab2 = tabs.Table2()
tab3 = tabs.Table3()


def tdlist_to_csv(data: list) -> str:
    """Converts a list of lists to a CSV string."""
    return "\n".join([",".join(x) for x in data])


def mercells(table: Table, cells: List[Tuple[int]]) -> None:
    if len(cells) > 1:
        try:
            cells[1] = table.cell(*cells[0]).merge(table.cell(*cells[1]))
        except Exception:
            return cells
        mercells(table, cells[1:])
    return cells


def tdlist_to_word(doc: Document, data: list) -> Table:
    l_rows = len(data)
    l_cols = max(map(len, data))
    table = doc.add_table(l_rows, l_cols, "Table Grid")

    p_el = ""
    # fill the table
    for i, row_data in enumerate(data):
        row = table.rows[i].cells
        for j, el in enumerate(row_data):
            # if p_el == el:
            #     mercells(table, [(i, j), (i, j - 1)])
            row[j].text = el
            p_el = el

    # merge cells verticaly
    # for i in range(l_rows):
    #     for j in range(1, l_cols):
    #         try:
    #             el = data[j][i]
    #             if data[j-1][i] == el:
    #                 mercells(table, [(j - 1, i), (j, i)])
    #             row = table.rows[j].cells
    #             row[i].text = el
    #         except IndexError:
    #             continue

    return table


if __name__ == "__main__":
    # Write the tables to CSV files
    # with open("./out/table1.csv", "w+", encoding="UTF-8") as f:
    #     f.write(tdlist_to_csv(table1.table(CODES, RESULTS)))\

    tdlist_to_word(DOC, tab1.table(CODES, RESULTS))
    DOC.add_page_break()
    for l, data in CODES.items():
        if l in "yzwl":
            continue
        args = [x for x in data.keys() if x != "_" and x != "func" and x != "zones"]
        args = (
            l,
            args,
        )

        tdlist_to_word(DOC, tab2.table(CODES, RESULTS, *args))
        DOC.add_page_break()

        for cat in ("13-15", "16-17", "18-19"):
            tdlist_to_word(DOC, tab3.table(CODES, RESULTS, cat, *args))
            DOC.add_page_break()

    DOC.save("./result.docx")
